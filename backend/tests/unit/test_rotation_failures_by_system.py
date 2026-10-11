"""Reporte `rotation-failures-by-system` (Excel por tipo de sistema)."""

from __future__ import annotations

import io
import time

from openpyxl import load_workbook

from app.reports import REGISTRY
from app.reports.base import ReportResult, result_to_dict
from app.reports.rotation_failures_by_system import (
    COLUMNS,
    UNCLASSIFIED,
    compute_by_system,
    platform_system_types,
    report,
)
from app.reports.xlsx import build_rotation_failures_xlsx, sheet_title

NOW = int(time.time())
DAY = 86400


def _acc(acc_id, platform, *, auto=True, status="success", age_days=1, kind="password", **kw):
    sm = {
        "automaticManagementEnabled": auto,
        "status": status,
        "lastModifiedTime": NOW - age_days * DAY,
    }
    sm.update(kw.pop("sm", {}))
    return {
        "id": acc_id,
        "name": f"acc{acc_id}",
        "userName": kw.pop("user", f"u{acc_id}"),
        "address": "10.0.0.1",
        "platformId": platform,
        "safeName": kw.pop("safe", "S1"),
        "secretType": kind,
        "secretManagement": sm,
    }


PLATFORMS = [
    {"general": {"id": "WinDomain", "systemType": "Windows"}},
    {"general": {"id": "WinServerLocal", "systemType": "Windows"}},
    {"general": {"id": "UnixSSH", "systemType": "*NIX"}},
    {"general": {"id": "Oracle", "systemType": "Database"}},
]

ACCOUNTS = [
    # Windows: 5 cuentas consideradas, 3 con problemas
    _acc("w-ok", "WinDomain"),
    _acc("w-fail", "WinDomain", status="failure"),
    _acc("w-over", "WinServerLocal", age_days=100),  # vencida, umbral por defecto 30
    _acc(
        "w-cpmoff",
        "WinDomain",
        auto=False,
        age_days=2,
        sm={"manualManagementReason": "(CPM)MaxRetries"},
    ),
    # deshabilitada por el cliente y vencida: cuenta en el total, no es problema
    _acc("w-manual", "WinDomain", auto=False, age_days=300, sm={"manualManagementReason": "x"}),
    # Unix: 2 cuentas, ninguna con problemas
    _acc("u-ok", "UnixSSH"),
    _acc("u-ok2", "UnixSSH"),
    # plataforma desconocida
    _acc("x-fail", "Rara", status="failure"),
    # excluidas siempre
    _acc("int", "WinDomain", status="failure", user="PSMConnect"),
    _acc("key", "UnixSSH", kind="key", age_days=400),
]


def _compute():
    return compute_by_system(
        ACCOUNTS,
        system_types=platform_system_types(PLATFORMS),
        platform_days={},
        cpm_disabled_ids=set(),
        default_stale_days=30,
    )


def test_registrado_con_excel():
    assert REGISTRY["rotation-failures-by-system"] is report
    assert report.heavy is True
    assert report.catalog()["formats"] == ["csv", "json", "xlsx"]
    assert REGISTRY["non-rotating-accounts"].catalog()["formats"] == ["csv", "json"]


def test_agrupa_por_sistema_y_calcula_el_porcentaje():
    rows, meta = _compute()
    by = {s["system"]: s for s in meta["by_system"]}
    assert by["Windows"] == {
        "system": "Windows",
        "total": 5,
        "cpm_errors": 2,
        "overdue": 1,
        "failing": 3,
        "pct": 60.0,
    }
    assert by["*NIX"]["failing"] == 0 and by["*NIX"]["total"] == 2  # aparece igual
    assert by[UNCLASSIFIED]["failing"] == 1
    # más problemas primero
    assert meta["by_system"][0]["system"] == "Windows"
    assert meta["cuentas_consideradas"] == 8
    assert meta["con_problemas"] == 4
    assert meta["porcentaje"] == 50.0
    assert meta["excluidas_internas_cyberark"] == 1
    assert meta["excluidas_llaves_ssh"] == 1

    problems = {r["id"]: r["problem"] for r in rows}
    assert problems == {
        "w-fail": "Error del CPM",
        "w-cpmoff": "Deshabilitada por el CPM",
        "w-over": "Vencida (sin error)",
        "x-fail": "Error del CPM",
    }
    assert {r["id"] for r in rows if r["systemType"] == "Windows"} == {
        "w-fail",
        "w-cpmoff",
        "w-over",
    }


def test_deshabilitada_por_el_cpm_segun_el_saved_filter():
    acc = _acc("a", "WinDomain", auto=False, age_days=2)
    rows, _ = compute_by_system(
        [acc],
        system_types=platform_system_types(PLATFORMS),
        platform_days={},
        cpm_disabled_ids={"a"},
        default_stale_days=30,
    )
    assert rows[0]["problem"] == "Deshabilitada por el CPM"


def test_umbral_de_la_plataforma():
    acc = _acc("a", "WinDomain", age_days=10)
    kw = {
        "system_types": platform_system_types(PLATFORMS),
        "cpm_disabled_ids": set(),
        "default_stale_days": 30,
    }
    assert compute_by_system([acc], platform_days={}, **kw)[0] == []
    rows, _ = compute_by_system([acc], platform_days={"WinDomain": 7}, **kw)
    assert rows[0]["problem"] == "Vencida (sin error)"


def test_nombre_de_hoja_valido_y_unico():
    used: set[str] = set()
    assert sheet_title("*NIX", used) == "Unix"
    assert sheet_title("Network/Device: [core]?", used) == "Network Device   core"
    assert sheet_title("Portada", used) == "Portada (2)"
    largo = "x" * 40
    assert sheet_title(largo, used) == "x" * 31
    assert sheet_title(largo, used) == "x" * 27 + " (2)"


def test_excel_portada_y_una_hoja_por_sistema_con_problemas():
    rows, meta = _compute()
    data = result_to_dict(ReportResult(columns=COLUMNS, rows=rows, meta=meta))
    data["generated_at"] = "2026-10-07T15:00:00+00:00"
    wb = load_workbook(io.BytesIO(build_rotation_failures_xlsx(data, "Globex")))

    # *NIX no tiene problemas: no tiene pestaña
    assert wb.sheetnames == ["Portada", "Windows", UNCLASSIFIED]

    portada = wb["Portada"]
    assert portada["B2"].value == "Globex"
    assert portada["A6"].value == "Sistema"
    filas = {
        portada.cell(row=r, column=1).value: [
            portada.cell(row=r, column=c).value for c in range(2, 7)
        ]
        for r in range(7, 12)
    }
    assert filas["Windows"] == [5, 2, 1, 3, 0.6]
    assert filas["*NIX"] == [2, 0, 0, 0, 0]
    assert filas["Total"] == [8, 3, 1, 4, 0.5]
    assert portada["F7"].number_format == "0.0%"
    assert portada["A7"].hyperlink.location == "'Windows'!A1"

    win = wb["Windows"]
    encabezados = [c.value for c in win[1]]
    assert "Sistema" not in encabezados  # la pestaña ya es el sistema
    assert encabezados[:3] == ["Account ID", "Nombre", "Usuario"]
    assert win.max_row == 4  # encabezado + 3 cuentas
    assert win.freeze_panes == "A2"
    assert win.auto_filter.ref == f"A1:{win.cell(row=1, column=win.max_column).column_letter}4"
