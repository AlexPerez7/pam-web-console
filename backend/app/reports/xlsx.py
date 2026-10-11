"""Exportación a Excel de reportes del catálogo que la declaran
(`Report.xlsx`). Hoy sólo `rotation-failures-by-system`: Portada con el %
de cuentas con problemas de rotación por sistema + una pestaña por sistema.
Ver `docs/specs/done/rotation-failures-by-system.md`.

Recibe el resultado ya serializado (`result_to_dict` + `generated_at`), el
mismo que queda cacheado por el job, así el Excel sale sin volver a
consultar CyberArk.
"""

from __future__ import annotations

import io
import re
from datetime import UTC, datetime
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from app.config import get_settings

NAVY = "002142"
_HEADER_FILL = PatternFill("solid", fgColor=NAVY)
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TOTAL_FILL = PatternFill("solid", fgColor="E8EDF3")
_THIN = Side(style="thin", color="C9D2DC")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_COVER = "Portada"
# ancho máximo de una columna de detalle (los motivos de error son largos)
_MAX_WIDTH = 60


def sheet_title(system: str, used: set[str]) -> str:
    """Nombre de hoja válido para Excel y único en el libro: sin `[]:*?/\\`,
    máx. 31 caracteres, `*NIX` -> `Unix`, distinto de "Portada"."""
    base = "Unix" if system.strip().upper() == "*NIX" else system
    base = _INVALID_SHEET_CHARS.sub(" ", base).strip().strip("'") or "Sin nombre"
    base = base[:31]
    title, n = base, 2
    while title.lower() in used or title.lower() == _COVER.lower():
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def _local(dt: datetime) -> datetime:
    """A la zona horaria de PAM Web Console y sin tzinfo (Excel no guarda zonas)."""
    tz = ZoneInfo(get_settings().alerts_timezone)
    return dt.astimezone(tz).replace(tzinfo=None)


def _epoch_to_local(value: Any) -> datetime | None:
    try:
        ts = int(value)
    except (TypeError, ValueError):
        return None
    if ts <= 0:
        return None
    return _local(datetime.fromtimestamp(ts, tz=UTC))


def _cell_value(kind: str, value: Any) -> Any:
    if value is None or value == "":
        return None
    if kind == "date":
        return _epoch_to_local(value)
    if kind == "bool":
        return "Sí" if value else "No"
    return value


def _header(ws: Worksheet, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.border = _BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _autowidth(ws: Worksheet, first_row: int, min_width: int = 10) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=first_row):
        for c in row:
            if c.value is None:
                continue
            text: Any = (
                c.value.strftime("%Y-%m-%d %H:%M") if isinstance(c.value, datetime) else c.value
            )
            widths[c.column] = max(widths.get(c.column, 0), len(str(text)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(w + 2, min_width), _MAX_WIDTH)


def _cover(ws: Worksheet, data: dict[str, Any], tenant_name: str, sheets: dict[str, str]) -> None:
    meta = data.get("meta") or {}
    generated = data.get("generated_at")
    when = _local(datetime.fromisoformat(generated)) if generated else _local(datetime.now(UTC))

    ws["A1"] = "Problemas de rotación de credenciales por sistema"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = "Cliente"
    ws["B2"] = tenant_name
    ws["A3"] = "Generado"
    ws["B3"] = when
    ws["B3"].number_format = "yyyy-mm-dd hh:mm"
    ws["B3"].alignment = Alignment(horizontal="left")
    ws["A4"] = "Criterio"
    ws["B4"] = (
        "Cuentas con error del CPM (o deshabilitadas por el CPM tras fallar) y cuentas "
        "gestionadas que superaron el plazo de rotación de su plataforma. El porcentaje se "
        "calcula sobre todas las cuentas del sistema."
    )
    ws["B4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B4:F4")
    ws.row_dimensions[4].height = 45
    for ref in ("A2", "A3", "A4"):
        ws[ref].font = Font(bold=True)
        ws[ref].alignment = Alignment(vertical="top")

    head = 6
    _header(
        ws,
        head,
        [
            "Sistema",
            "Cuentas",
            "Con error del CPM",
            "Vencidas (sin error)",
            "Con problemas",
            "% con problemas",
        ],
    )
    row = head
    for s in meta.get("by_system") or []:
        row += 1
        name = ws.cell(row=row, column=1, value=s["system"])
        if s["system"] in sheets:  # enlace a su pestaña
            name.hyperlink = Hyperlink(ref=name.coordinate, location=f"'{sheets[s['system']]}'!A1")
            name.font = Font(color="0B5CAD", underline="single")
        values = [s["total"], s["cpm_errors"], s["overdue"], s["failing"], s["pct"] / 100]
        for col, v in enumerate(values, start=2):
            ws.cell(row=row, column=col, value=v)
        ws.cell(row=row, column=6).number_format = "0.0%"
        for col in range(1, 7):
            ws.cell(row=row, column=col).border = _BORDER

    first, last = head + 1, row
    row += 1
    ws.cell(row=row, column=1, value="Total")
    total = meta.get("cuentas_consideradas", 0)
    failing = meta.get("con_problemas", 0)
    by_system = meta.get("by_system") or []
    totals = [
        total,
        sum(s["cpm_errors"] for s in by_system),
        sum(s["overdue"] for s in by_system),
        failing,
        (meta.get("porcentaje") or 0) / 100,
    ]
    for col, v in enumerate(totals, start=2):
        ws.cell(row=row, column=col, value=v)
    ws.cell(row=row, column=6).number_format = "0.0%"
    for col in range(1, 7):
        c = ws.cell(row=row, column=col)
        c.font = Font(bold=True)
        c.fill = _TOTAL_FILL
        c.border = _BORDER

    if last >= first:
        ws.conditional_formatting.add(
            f"F{first}:F{last}",
            DataBarRule(
                start_type="num", start_value=0, end_type="num", end_value=1, color="C0392B"
            ),
        )

    notes = [
        f"Se excluyeron {meta.get('excluidas_internas_cyberark', 0)} cuentas internas de "
        "CyberArk (componentes PSM/CPM/PVWA) y "
        f"{meta.get('excluidas_llaves_ssh', 0)} llaves SSH: no son credenciales que deban rotar.",
        "Las cuentas con gestión automática deshabilitada a propósito cuentan en el total, "
        "pero no como problema.",
    ]
    if meta.get("ultimo_error_sin_datos"):
        notes.append(
            f"No se pudo consultar el último error de {meta['ultimo_error_sin_datos']} cuentas."
        )
    row += 2
    for note in notes:
        ws.cell(row=row, column=1, value=note).font = Font(italic=True, color="555555")
        row += 1

    ws.column_dimensions["A"].width = 24
    for letter in "BCDEF":
        ws.column_dimensions[letter].width = 18
    ws.sheet_view.showGridLines = False


def _detail(ws: Worksheet, columns: list[dict[str, Any]], rows: list[dict[str, Any]]) -> None:
    _header(ws, 1, [c["label"] for c in columns])
    for r, row in enumerate(rows, start=2):
        for col, c in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=col, value=_cell_value(c["kind"], row.get(c["key"])))
            if c["kind"] == "date":
                cell.number_format = "yyyy-mm-dd hh:mm"
    ws.freeze_panes = "A2"
    if not columns:
        return
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 1)}"
    _autowidth(ws, 1)


def build_rotation_failures_xlsx(data: dict[str, Any], tenant_name: str) -> bytes:
    """Portada + una pestaña por sistema con al menos una cuenta con
    problemas, en el orden de la Portada (más problemas primero)."""
    columns = [c for c in data.get("columns") or [] if c["key"] != "systemType"]
    by_system = (data.get("meta") or {}).get("by_system") or []
    rows_by_system: dict[str, list[dict[str, Any]]] = {}
    for row in data.get("rows") or []:
        rows_by_system.setdefault(row.get("systemType") or "", []).append(row)

    wb = Workbook()
    cover = wb.worksheets[0]  # un Workbook nuevo siempre trae una hoja
    cover.title = _COVER

    used: set[str] = set()
    sheets: dict[str, str] = {}
    for s in by_system:
        rows = rows_by_system.get(s["system"])
        if not rows:
            continue
        title = sheet_title(s["system"], used)
        sheets[s["system"]] = title
        _detail(wb.create_sheet(title), columns, rows)

    _cover(cover, data, tenant_name, sheets)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
